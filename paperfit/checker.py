# -*- coding: utf-8 -*-
"""PaperFit 核心：按规则逐条体检 Word 论文格式。"""
import re
from docx import Document
from docx.oxml.ns import qn

CN_ALIGN = {
    "center": "居中", "left": "左对齐", "start": "左对齐",
    "right": "右对齐", "justify": "两端对齐", "both": "两端对齐", None: "继承",
}

def _val(el):
    if el is None:
        return None
    v = el.get(qn("w:val"))
    if v is None:
        return True
    return v not in ("0", "false")

def run_style(run):
    rPr = run._element.find(qn("w:rPr"))
    cn = en = None
    size = None
    bold = False
    if rPr is not None:
        rf = rPr.find(qn("w:rFonts"))
        if rf is not None:
            cn = rf.get(qn("w:eastAsia"))
            en = rf.get(qn("w:ascii"))
        z = rPr.find(qn("w:sz"))
        if z is not None:
            try:
                size = int(z.get(qn("w:val"))) / 2.0
            except Exception:
                pass
        b = rPr.find(qn("w:b"))
        bold = _val(b) is True
    return {"cn": cn, "en": en, "size": size, "bold": bold}

def para_style(p):
    pPr = p._p.find(qn("w:pPr"))
    align = None
    indent_first = 0
    if pPr is not None:
        jc = pPr.find(qn("w:jc"))
        if jc is not None:
            align = jc.get(qn("w:val"))
        ind = pPr.find(qn("w:ind"))
        if ind is not None:
            for k, v in ind.attrib.items():
                if k.endswith("}firstLine") or k.endswith("}firstLineChars"):
                    try:
                        indent_first = int(v)
                    except Exception:
                        pass
    return {"align": align, "indent_first": indent_first}

def classify(paragraphs):
    """把段落分角色。返回 [(index, role, text, para)]"""
    out = []
    seen_title = False
    in_refs = False
    for i, p in enumerate(paragraphs):
        t = p.text.strip()
        if not t:
            continue
        if not seen_title:
            seen_title = True
            out.append((i, "title", t, p)); continue
        if len(out) == 1:
            out.append((i, "author", t, p)); continue
        if len(out) == 2 and (t.startswith("（") or t.startswith("(") or "大学" in t or "学院" in t):
            out.append((i, "org", t, p)); continue
        if t.startswith("摘要"):
            out.append((i, "abstract", t, p)); continue
        if t.startswith("关键词") or t.startswith("关键字"):
            out.append((i, "keywords", t, p)); continue
        if t.startswith("参考文献"):
            in_refs = True
            continue
        if in_refs or re.match(r"^\[\d+\]", t):
            out.append((i, "ref", t, p)); continue
        if re.match(r"^\d+(\.\d+)*[\s　]+", t) or re.match(r"^[一二三四五六七八九十]+[、\s]", t):
            if re.match(r"^\d+\.\d+", t):
                out.append((i, "h2", t, p))
            else:
                out.append((i, "h1", t, p))
            continue
        out.append((i, "body", t, p))
    return out

def check(path, rules):
    doc = Document(path)
    items = classify(doc.paragraphs)
    issues, passed = [], []
    stats = {}

    # 页面
    sec = doc.sections[0]
    page = rules.get("page", {})
    page_ok = True
    if page:
        w_cm = round(sec.page_width.cm, 2)
        h_cm = round(sec.page_height.cm, 2)
        if abs(w_cm - page.get("width_cm", w_cm)) > 0.2 or abs(h_cm - page.get("height_cm", h_cm)) > 0.2:
            page_ok = False
            issues.append({"role": "页面", "expect": f"{page['width_cm']}×{page['height_cm']} cm",
                           "actual": f"{w_cm}×{h_cm} cm", "where": "文档设置", "level": "error"})
    if page_ok:
        passed.append("页面尺寸")

    role_rules = rules.get("roles", {})
    for role, rule in role_rules.items():
        group = [it for it in items if it[1] == role]
        if not group:
            continue
        stats[role] = len(group)
        bad_count = 0
        for idx, r, text, p in group:
            if not p.runs:
                continue
            st = run_style(p.runs[0])
            ps = para_style(p)
            exp_desc, act_desc = [], []
            ok = True
            if rule.get("cn_font") and st["cn"] and st["cn"] != rule["cn_font"]:
                ok = False
                exp_desc.append(f"中文{rule['cn_font']}")
                act_desc.append(f"中文{st['cn']}")
            if rule.get("size_pt") and st["size"] and abs(st["size"] - rule["size_pt"]) > 0.05:
                ok = False
                exp_desc.append(f"{rule['size_pt']}pt")
                act_desc.append(f"{st['size']}pt")
            if "bold" in rule and st["bold"] != rule["bold"]:
                ok = False
                exp_desc.append("加粗" if rule["bold"] else "不加粗")
                act_desc.append("加粗" if st["bold"] else "不加粗")
            if rule.get("align"):
                want = rule["align"]
                cur = ps["align"] or "start"
                norm = {"both": "justify", "start": "left"}
                if norm.get(cur, cur) != norm.get(want, want):
                    ok = False
                    exp_desc.append(CN_ALIGN.get(want, want))
                    act_desc.append(CN_ALIGN.get(cur, cur))
            if rule.get("indent_chars") and role in ("body", "abstract"):
                need = rule["indent_chars"] * (rule.get("size_pt") or 10.5) * 20
                if abs(ps["indent_first"] - need) > 60 and ps["indent_first"] < 200:
                    ok = False
                    exp_desc.append(f"首行缩进{rule['indent_chars']}字符")
                    act_desc.append("无缩进" if ps["indent_first"] == 0 else f"缩进{ps['indent_first']:.0f}twip")
            if not ok:
                bad_count += 1
                if bad_count <= 3:
                    issues.append({"role": rule.get("label", role), "expect": "、".join(exp_desc),
                                   "actual": "、".join(act_desc), "where": f"第{idx+1}段: {text[:24]}",
                                   "level": "error"})
        if bad_count == 0:
            passed.append(rule.get("label", role))
        elif bad_count > 3:
            issues.append({"role": rule.get("label", role), "expect": "（同上）", "actual": f"共 {bad_count} 处不符",
                           "where": "多处", "level": "error"})

    # 参考文献数量
    rr = rules.get("reference", {})
    refs = [it for it in items if it[1] == "ref"]
    if rr.get("min_count") and len(refs) < rr["min_count"]:
        issues.append({"role": "参考文献数量", "expect": f"≥{rr['min_count']} 篇",
                       "actual": f"{len(refs)} 篇", "where": "文末", "level": "error"})
    else:
        passed.append("参考文献数量")

    total = len(passed) + len(issues)
    score = round(len(passed) / total * 100) if total else 0
    return {
        "file": path.split("\\")[-1],
        "rules": rules.get("name", "未命名规则"),
        "score": score,
        "passed": passed,
        "issues": issues,
        "stats": stats,
        "paragraphs": len(doc.paragraphs),
    }
